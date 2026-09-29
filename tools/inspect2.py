# 只读：聚焦四个「油耗可视化」sheet 的字段位置与图片列
import openpyxl
from openpyxl.utils import get_column_letter

PATH = r"C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm"
wb = openpyxl.load_workbook(PATH, data_only=True, keep_links=False)

targets = [n for n in wb.sheetnames if "可视化" in n and "尿素" not in n]
print("目标 sheet:", targets)

for name in targets:
    ws = wb[name]
    imgs = getattr(ws, "_images", [])
    print(f"\n===== {name!r}  行={ws.max_row} 列={ws.max_column} 现有浮动图片={len(imgs)}")
    for r in range(1, 14):
        cells = []
        for c in range(1, 13):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip() != "":
                cells.append(f"{get_column_letter(c)}{r}={str(v)[:30]}")
        if cells:
            print("   " + " | ".join(cells))
    for im in imgs:
        try:
            a = im.anchor._from
            print(f"   [图片] 锚点行{a.row}(+1={a.row+1}) 列{a.col}({get_column_letter(a.col+1)})")
        except Exception as e:
            print("   [图片] 锚点解析失败", e)
    # I 列（平均时速值右侧）逐行看有无内容
    col_i = []
    for r in range(1, min(ws.max_row, 40) + 1):
        v = ws.cell(row=r, column=9).value
        if v is not None and str(v).strip() != "":
            col_i.append(f"I{r}={str(v)[:30]}")
    print("   I列内容:", col_i if col_i else "(空)")
    # 合并单元格
    try:
        mc = [str(m) for m in ws.merged_cells.ranges][:12]
        print("   合并单元格:", mc)
    except Exception:
        pass
