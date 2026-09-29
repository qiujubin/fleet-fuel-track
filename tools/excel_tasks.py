# 从四个「油耗可视化」sheet 生成轨迹截图任务（只读，不写表）
# 规则：无车牌 / 时间非日期 / 跨度 > 10 天 → 跳过
import json
import datetime
import openpyxl

PATH = r"C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm"
SHEETS = ["干线油耗可视化", "华北油耗可视化 ", "华东油耗可视化 ", "华南油耗可视化"]
MAX_DAYS = 10

wb = openpyxl.load_workbook(PATH, data_only=True, keep_links=False)
tasks, skipped = [], []

for name in SHEETS:
    if name not in wb.sheetnames:
        skipped.append((name, "sheet 不存在"))
        continue
    ws = wb[name]
    plate = ws["C2"].value
    start, end = ws["C4"].value, ws["C5"].value
    reason = None
    if not plate or not str(plate).strip():
        reason = "无车牌"
    elif not isinstance(start, datetime.datetime) or not isinstance(end, datetime.datetime):
        reason = f"时间非日期(start={start!r}, end={end!r})"
    elif (end - start).total_seconds() > MAX_DAYS * 86400:
        reason = f"跨度 {(end - start).total_seconds() / 86400:.1f} 天 > {MAX_DAYS} 天"
    if reason:
        skipped.append((name, reason))
        continue
    tasks.append({
        "plate": str(plate).strip(),
        "start": start.strftime("%Y-%m-%d %H:%M:%S"),
        "end": end.strftime("%Y-%m-%d %H:%M:%S"),
        "file": f"track_{str(plate).strip()}.png",
        "sheet": name,
    })

print(f"待处理 {len(tasks)} 项：")
for t in tasks:
    print(f"  [{t['sheet'].strip()}] {t['plate']}  {t['start']} ~ {t['end']}  -> {t['file']}")
print(f"\n跳过 {len(skipped)} 项：")
for n, r in skipped:
    print(f"  [{n.strip()}] {r}")

with open("tasks_from_excel.json", "w", encoding="utf-8") as f:
    json.dump(tasks, f, ensure_ascii=False, indent=2)
print("\n已生成 tasks_from_excel.json")
