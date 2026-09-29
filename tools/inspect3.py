# 只读：探查用户指定的四个区域 sheet：干线油耗 / 华北油耗 / 华东油耗 / 华南油耗
import openpyxl
from openpyxl.utils import get_column_letter

PATH = r"C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm"
wb = openpyxl.load_workbook(PATH, data_only=True, keep_links=False)
print("sheet 总数:", len(wb.sheetnames))

targets = ["干线油耗", "华北油耗", "华东油耗", "华南油耗"]
for t in targets:
    if t not in wb.sheetnames:
        print(f"\n!!! 找不到 sheet: {t}")
        continue
    ws = wb[t]
    print(f"\n===== {t}  行={ws.max_row} 列={ws.max_column} 图片={len(getattr(ws,'_images',[]))}")
    for r in range(1, min(ws.max_row, 40) + 1):
        cells = []
        for c in range(1, min(ws.max_column, 24) + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip() != "":
                cells.append(f"{get_column_letter(c)}={str(v)[:22]}")
        if cells:
            print(f"  第{r}行: " + " | ".join(cells))
    try:
        print("  合并:", [str(m) for m in ws.merged_cells.ranges][:10])
    except Exception:
        pass
